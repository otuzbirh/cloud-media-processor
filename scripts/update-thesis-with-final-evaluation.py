#!/usr/bin/env python3
"""Insert the final practical evaluation into the existing master's thesis DOCX."""

from __future__ import annotations

import argparse
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = PROJECT_ROOT / "output/documents/Master_rad_UNZE_PTF_sa_prakticnom_evaluacijom.docx"
DEFAULT_OUTPUT = PROJECT_ROOT / "output/documents/Master_rad_UNZE_PTF_sa_finalnom_evaluacijom.docx"
DEFAULT_CHAPTER = PROJECT_ROOT / "docs/practical-evaluation-draft.md"
FIGURE_DIR = PROJECT_ROOT / "analysis/final/thesis"

TABLE_CAPTIONS = [
    "Tabela 7.1. Eksperimentalne politike alokacije worker replika",
    "Tabela 7.2. Profili eksperimentalnog opterećenja",
    "Tabela 7.3. Rezultati lokalnog pilot-testa",
    "Tabela 7.4. Aktivirani kapacitet i efikasnost workera",
    "Tabela 7.5. Vrijeme obrta, čekanje i throughput",
    "Tabela 7.6. Dinamika automatskog skaliranja",
    "Tabela 7.7. Eksperimentalna ocjena istraživačkih hipoteza",
]

FIGURES_AFTER_TABLE = {
    4: ("worker-minute-po-profilu.png", "Slika 7.1. Aktivirani worker-kapacitet po profilu opterećenja", 6.15, False),
    5: ("vrijeme-obrta-po-profilu.png", "Slika 7.2. Prosječno vrijeme obrta po profilu opterećenja", 6.15, False),
    6: ("dinamika-high.png", "Slika 7.3. Dinamika sistema tokom profila high", 5.95, True),
}

FIGURES_AFTER_PARAGRAPH = {
    "KEDA je u profilu `high`": ("dinamika-ramp.png", "Slika 7.4. Dinamika sistema tokom profila ramp", 5.95, True),
    "U profilu `spike` oba skalera": ("dinamika-spike.png", "Slika 7.5. Dinamika sistema tokom profila spike", 5.95, True),
}

TABLE_WIDTHS = {
    1: [0.62, 1.45, 4.18],
    2: [0.85, 1.85, 3.55],
    3: [3.20, 1.52, 1.52],
    4: [0.55, 0.83, 1.45, 0.88, 0.86, 1.68],
    5: [0.55, 0.83, 1.42, 0.95, 1.25, 1.25],
    6: [0.55, 0.88, 1.52, 1.04, 0.80, 1.08],
    7: [0.72, 1.55, 3.98],
}

PARAGRAPH_REPLACEMENTS = {
    "Upravljanje resursima i njihova alokacija određuju performanse": (
        "Upravljanje resursima i njihova alokacija određuju performanse, trošak i efikasnost cloud usluga. "
        "Cilj rada je sistematizovati ključne politike upravljanja resursima i eksperimentalno ispitati odnos aktiviranog kapaciteta i kvaliteta usluge pri statičkoj alokaciji, CPU HPA i KEDA skaliranju."
    ),
    "Istraživanje povezuje analitičku, deskriptivnu": (
        "Istraživanje povezuje analitičku, deskriptivnu, sintetičku i komparativnu metodu sa eksperimentalnom metodom. "
        "Razvijena je kontejnerizovana aplikacija za asinhronu batch obradu slika, mjerni postupak zasnovan na eksperimentalnim sesijama i automatizovana matrica opterećenja. "
        "Konačna evaluacija obuhvatila je statičku alokaciju, CPU HPA i KEDA skaliranje prema dužini reda, četiri profila opterećenja i pet ponavljanja svake kombinacije."
    ),
    "Pri malom opterećenju red nije formiran": (
        "U 60 validnih izvođenja poslano je i završeno 67.954 posla bez greške. Dinamičke politike smanjile su aktivirani worker-kapacitet u svim profilima, od 14,7% do 83,4% u odnosu na statičku alokaciju. "
        "Najveća ušteda ostvarena je KEDA politikom pri malom opterećenju, bez pogoršanja vremena obrta. Pri trajnom i naglom visokom opterećenju smanjenje kapaciteta povećalo je čekanje, što potvrđuje kompromis između efikasnosti resursa i repne latencije."
    ),
    "The research combines analytical, descriptive": (
        "The research combines analytical, descriptive, synthetic, comparative, and experimental methods. A containerized application for asynchronous batch image processing, a session-based measurement process, and an automated workload matrix were developed. "
        "The final evaluation covered static allocation, CPU-based horizontal scaling, queue-based KEDA scaling, four workload profiles, and five repetitions of every combination."
    ),
    "Resource management and allocation determine the performance": (
        "Resource management and allocation determine the performance, cost, and efficiency of cloud services. "
        "This thesis systematizes the principal resource-management policies and experimentally examines the relationship between activated capacity and quality of service under static allocation, CPU-based HPA, and KEDA scaling."
    ),
    "Under low load, no queue formed": (
        "Across 60 valid runs, 67,954 jobs were submitted and completed without failure. The dynamic policies reduced activated worker capacity in every profile by 14.7% to 83.4% relative to static allocation. "
        "Queue-based KEDA achieved the largest reduction under low load without degrading turnaround time. Under sustained and sudden high load, lower initial capacity increased queueing, demonstrating the trade-off between resource efficiency and tail latency."
    ),
    "U radu su primijenjene analitička, sintetička": (
        "U radu su primijenjene analitička, sintetička, induktivna, komparativna, deskriptivna i eksperimentalna metoda. Teorijska analiza dopunjena je lokalnim pilotom i konačnom Kubernetes evaluacijom sa 60 strogo validiranih izvođenja. "
        "Bilježeni su aktivirani kapacitet, protok, čekanje, vrijeme obrta, iskorištenost resursa, reakcija skaliranja i greške."
    ),
    "Rad najprije prikazuje teorijske osnove": (
        "Rad najprije prikazuje teorijske osnove cloud computinga i upravljanja resursima, zatim analizira politike alokacije, scheduling i dinamičko skaliranje. "
        "Nakon teorijske i komparativne analize prikazani su pilot mjernog postupka i konačno eksperimentalno poređenje statičke, CPU HPA i KEDA politike."
    ),
    "Induktivna metoda korištena je za izvođenje općih zaključaka": (
        "Induktivna metoda korištena je za izvođenje općih zaključaka iz analiziranih pristupa i opažanja pilota i konačne eksperimentalne matrice, uz jasno navedena ograničenja njihovog opsega."
    ),
    "Eksperimentalna metoda primijenjena je u lokalnom pilotu": (
        "Eksperimentalna metoda primijenjena je u lokalnom pilotu i kontrolisanoj Kubernetes matrici sa tri politike alokacije, četiri profila opterećenja i pet ponavljanja. "
        "Rezultati su analizirani deskriptivnim pokazateljima i vremenskim serijama. Teorijski dio zasnovan je na relevantnoj naučnoj i stručnoj literaturi."
    ),
    "Sedmo poglavlje prikazuje lokalnu pilot-evaluaciju": (
        "Sedmo poglavlje prikazuje lokalni pilot, eksperimentalno okruženje, strogu kontrolu valjanosti i konačno poređenje statičke, CPU HPA i KEDA politike. Zaključak objedinjuje teorijske i praktične nalaze, ograničenja i preporuke za izbor politike."
    ),
    "H1 je na osnovu teorijske i ilustrativne analize podržana uslovno": (
        "H1 je na osnovu teorijske i ilustrativne analize podržana uslovno. Praktična evaluacija u sedmom poglavlju provjerava ovaj mehanizam poređenjem statičke i dinamičkih politika pri istom opterećenju, uz odvojeno izvještavanje aktiviranog kapaciteta i kvaliteta usluge."
    ),
    "H2 je teorijski podržana uz uslov": (
        "H2 je teorijski podržana uz uslov da se efikasnost mjeri skupom pokazatelja. Veća iskorištenost nije poboljšanje ako povećava greške ili neprihvatljivo pogoršava vrijeme obrta. Sedmo poglavlje zato zajedno posmatra kapacitet, protok, čekanje, vrijeme obrta i stopu grešaka."
    ),
    "Ocjene u ovom poglavlju predstavljaju teorijsku osnovu": (
        "Ocjene u ovom poglavlju predstavljaju teorijsku osnovu za praktičnu evaluaciju. Konačna eksperimentalna ocjena hipoteza izvedena je u sedmom poglavlju i formulirana je uslovno, u granicama ispitane aplikacije i lokalnog Kubernetes okruženja."
    ),
    "Lokalni pilot pokazao je dvije krajnosti statičke alokacije": (
        "Lokalni pilot potvrdio je osjetljivost mjernog postupka, a konačna matrica omogućila poređenje politika. U 60 validnih izvođenja završeno je svih 67.954 poslova bez greške i bez odbačenih iteracija."
    ),
    "Pilot je također otkrio metodološku grešku": (
        "Pilot je otkrio metodološku grešku u početnom statusnom anketiranju i omogućio njeno uklanjanje prije konačnih mjerenja. Stroga analiza zatim je potvrdila pet ponavljanja svake od 12 kombinacija, jednu Git reviziju i potpune artefakte."
    ),
    "Teorijska analiza uslovno podržava obje hipoteze": (
        "Konačni rezultati parcijalno podržavaju H1 jer su dinamičke politike smanjile worker-minute za 14,7% do 83,4%, ali ne dokazuju stvarnu novčanu uštedu na fiksno alociranom hostu. H2 je podržana uslovno: efikasnost aktiviranih workera je povećana bez gubitka poslova, dok su zahtjevniji profili pokazali cijenu u vidu većeg čekanja i vremena obrta."
    ),
    "Sljedeći istraživački korak je izvođenje pune matrice": (
        "Za posmatranu aplikaciju nema univerzalno najbolje politike: KEDA najviše štedi kapacitet, CPU HPA bolje prati postepeni rast, a statička alokacija najbolje štiti latenciju. "
        "Dalji rad treba uključiti višestruke čvorove, unaprijed definisan SLO i proaktivno skaliranje."
    ),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--chapter", type=Path, default=DEFAULT_CHAPTER)
    parser.add_argument("--page-map", type=Path)
    return parser.parse_args()


def clean_text(value: str) -> str:
    return value.replace("−", "-").replace("–", "-").replace("—", "-").replace("→", "->")


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def find_paragraph(document: Document, predicate) -> object:
    for paragraph in document.paragraphs:
        if predicate(paragraph):
            return paragraph
    raise ValueError("Traženi paragraf nije pronađen.")


def add_inline(paragraph, value: str) -> None:
    value = clean_text(value).replace("**", "")
    chunks = value.split("`")
    for index, chunk in enumerate(chunks):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        if index % 2:
            run.font.name = "Courier New"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
            run.font.size = Pt(9.5)


def insert_paragraph_before(document: Document, anchor, text: str = "", style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    if text:
        add_inline(paragraph, text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def set_cell_margins(cell, *, top: int = 70, start: int = 110, bottom: int = 70, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def insert_table(document: Document, anchor, rows: list[list[str]], table_number: int):
    caption = insert_paragraph_before(document, anchor, TABLE_CAPTIONS[table_number - 1], "Table Caption")
    caption.paragraph_format.keep_with_next = True
    row_groups = [rows]
    if table_number == 5:
        # Keep the profile groups together and provide an explicit header on the
        # continuation page in renderers that ignore Word's repeat-header flag.
        row_groups = [rows[:7], [rows[0], *rows[7:]]]

    tables = []
    for group_index, group_rows in enumerate(row_groups):
        if group_index:
            continuation = insert_paragraph_before(
                document,
                anchor,
                f"{TABLE_CAPTIONS[table_number - 1]} (nastavak)",
                "Table Caption",
            )
            continuation.paragraph_format.keep_with_next = True
        table = document.add_table(rows=len(group_rows), cols=len(group_rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Normal Table"
        widths = TABLE_WIDTHS[table_number]
        for row_index, row_values in enumerate(group_rows):
            row = table.rows[row_index]
            prevent_row_split(row)
            if row_index == 0:
                mark_repeat_header(row)
            for column_index, value in enumerate(row_values):
                cell = row.cells[column_index]
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, start=180 if column_index == 0 else 110)
                set_cell_borders(cell)
                if row_index == 0:
                    shade_cell(cell, "D9E5F3")
                elif row_index % 2 == 0:
                    shade_cell(cell, "F5F8FC")
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                add_inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
                    run.font.size = Pt(8.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
        anchor._p.addprevious(table._tbl)
        tables.append(table)
    source = insert_paragraph_before(document, anchor, "Izvor: izrada autora na osnovu eksperimentalnih mjerenja.", "Source")
    source.paragraph_format.space_after = Pt(8)
    return tables[0]


def insert_figure(document: Document, anchor, spec: tuple[str, str, float, bool]) -> None:
    filename, caption_text, width, page_break = spec
    image_path = FIGURE_DIR / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.page_break_before = page_break
    run = paragraph.add_run()
    drawing = run.add_picture(str(image_path), width=Inches(width))
    doc_pr = drawing._inline.docPr
    doc_pr.set("descr", caption_text)
    anchor._p.addprevious(paragraph._p)
    caption = insert_paragraph_before(document, anchor, caption_text, "Caption")
    caption.paragraph_format.keep_with_next = True
    source = insert_paragraph_before(document, anchor, "Izvor: izrada autora na osnovu eksperimentalnih mjerenja.", "Source")
    source.paragraph_format.space_after = Pt(8)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def markdown_headings(chapter_path: Path) -> list[tuple[int, str]]:
    headings = []
    for line in chapter_path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if match:
            headings.append((len(match.group(1)), clean_text(match.group(2))))
    return headings


def insert_chapter(document: Document, chapter_path: Path, anchor) -> None:
    lines = chapter_path.read_text(encoding="utf-8").splitlines()
    index = 0
    table_number = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = insert_paragraph_before(document, anchor, heading.group(2), f"Heading {level}")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            rows, index = parse_table(lines, index)
            table_number += 1
            insert_table(document, anchor, rows, table_number)
            if table_number in FIGURES_AFTER_TABLE:
                insert_figure(document, anchor, FIGURES_AFTER_TABLE[table_number])
            continue
        list_item = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if list_item:
            paragraph = insert_paragraph_before(document, anchor, f"{list_item.group(1)}. {list_item.group(2)}", "Normal")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.28)
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith("#") or candidate.startswith("|") or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph_text = " ".join(paragraph_lines)
        paragraph = insert_paragraph_before(document, anchor, paragraph_text, "Normal")
        paragraph.paragraph_format.widow_control = True
        for prefix, figure in FIGURES_AFTER_PARAGRAPH.items():
            if paragraph_text.startswith(prefix):
                insert_figure(document, anchor, figure)
    if table_number != len(TABLE_CAPTIONS):
        raise ValueError(f"Očekivano je {len(TABLE_CAPTIONS)} tabela, pronađeno {table_number}.")


def replace_paragraphs(document: Document) -> None:
    for prefix, replacement in PARAGRAPH_REPLACEMENTS.items():
        matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
        if len(matches) != 1:
            raise ValueError(f"Očekivan je jedan paragraf sa početkom {prefix!r}, pronađeno {len(matches)}.")
        paragraph = matches[0]
        paragraph.clear()
        add_inline(paragraph, replacement)


def replace_chapter(document: Document, chapter_path: Path) -> None:
    chapter = find_paragraph(document, lambda p: p.style.name == "Heading 1" and p.text.startswith("Praktična evaluacija"))
    conclusion = find_paragraph(document, lambda p: p.style.name == "Heading 1" and p.text.strip() == "ZAKLJUČAK")
    body = document._element.body
    children = list(body)
    start = children.index(chapter._p)
    end = children.index(conclusion._p)
    for element in children[start:end]:
        remove_element(element)
    insert_chapter(document, chapter_path, conclusion)


def split_front_entry(text: str) -> str:
    return text.split("\t", 1)[0].strip()


def apply_paragraph_template(paragraph, template) -> None:
    """Copy direct tab stops and spacing that are not carried by the style."""
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None:
        paragraph._p.remove(paragraph_properties)
    paragraph._p.insert(0, deepcopy(template._p.pPr))


def insert_front_entry(document: Document, anchor, label: str, page: str, template):
    paragraph = insert_paragraph_before(document, anchor)
    apply_paragraph_template(paragraph, template)
    paragraph.add_run(f"{label}\t{page}")
    return paragraph


def page_for(label: str, page_map: dict[str, str] | None, fallback: str = "??") -> str:
    if page_map and label in page_map:
        return str(page_map[label])
    return fallback


def update_front_matter(document: Document, chapter_path: Path, page_map: dict[str, str] | None) -> None:
    template_styles = ("TOC 1 Custom", "TOC 2 Custom", "TOC 3 Custom", "List Entry")
    templates = {
        style: find_paragraph(document, lambda p, style=style: p.style.name == style and "\t" in p.text)
        for style in template_styles
    }

    toc_conclusion = find_paragraph(document, lambda p: p.style.name.startswith("TOC") and split_front_entry(p.text) == "ZAKLJUČAK")
    for paragraph in list(document.paragraphs):
        if paragraph.style.name.startswith("TOC") and split_front_entry(paragraph.text).startswith("7."):
            remove_element(paragraph._p)
    for level, label in markdown_headings(chapter_path):
        style = f"TOC {level} Custom"
        insert_front_entry(document, toc_conclusion, label, page_for(label, page_map), templates[style])

    figures_heading = find_paragraph(document, lambda p: p.style.name == "Front Heading" and p.text.strip() == "POPIS TABELA")
    for paragraph in list(document.paragraphs):
        if paragraph.style.name == "List Entry" and split_front_entry(paragraph.text).startswith("Slika 7."):
            remove_element(paragraph._p)
    figure_labels = [spec[1] for spec in FIGURES_AFTER_TABLE.values()] + [spec[1] for spec in FIGURES_AFTER_PARAGRAPH.values()]
    for label in sorted(figure_labels, key=lambda value: int(re.search(r"7\.(\d+)", value).group(1))):
        insert_front_entry(document, figures_heading, label, page_for(label, page_map), templates["List Entry"])

    abbreviations_heading = find_paragraph(document, lambda p: p.style.name == "Front Heading" and p.text.strip() == "POPIS SKRAĆENICA")
    for paragraph in list(document.paragraphs):
        if paragraph.style.name == "List Entry" and split_front_entry(paragraph.text).startswith("Tabela 7."):
            remove_element(paragraph._p)
    for label in TABLE_CAPTIONS:
        insert_front_entry(document, abbreviations_heading, label, page_for(label, page_map), templates["List Entry"])

    if page_map:
        for paragraph in document.paragraphs:
            if paragraph.style.name.startswith("TOC") or paragraph.style.name == "List Entry":
                label = split_front_entry(paragraph.text)
                if label in page_map:
                    paragraph.text = f"{label}\t{page_map[label]}"
                    paragraph.style = paragraph.style


def enable_field_updates(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    args = parse_args()
    page_map = json.loads(args.page_map.read_text(encoding="utf-8")) if args.page_map else None
    document = Document(args.source)
    replace_paragraphs(document)
    replace_chapter(document, args.chapter)
    update_front_matter(document, args.chapter, page_map)
    enable_field_updates(document)
    document.core_properties.title = "Master rad - praktična evaluacija politika alokacije resursa"
    document.core_properties.subject = "Konačna Kubernetes evaluacija statičke alokacije, CPU HPA i KEDA skaliranja"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
